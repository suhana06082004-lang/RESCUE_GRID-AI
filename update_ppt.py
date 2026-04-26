import os
from pptx import Presentation
from pptx.util import Pt
import sys

ppt_path = os.path.abspath("Solution_Challenge_RescueGrid.pptx")
print(f"Reading from: {ppt_path}")
if not os.path.exists(ppt_path):
    print("File does not exist")
    sys.exit(1)

try:
    ppt = Presentation(ppt_path)
except Exception as e:
    print(f"Error loading PPTX: {e}")
    sys.exit(1)

slide2 = ppt.slides[1]

option3 = "Modern emergency management faces a dual point of failure during severe crises: fragile cloud infrastructure that collapses under blackout conditions, and human dispatchers who become overwhelmed by chaotic, multi-lingual distress signals. There is a critical technological gap in bridging raw human panic at the edge of a network, and instantly synthesizing it into prioritized, actionable deployment logistics for first responders without depending on vulnerable backend servers."

for shape in slide2.shapes:
    if hasattr(shape, "text") and "Team Details" in shape.text:
        shape.text = f"Team Details\n\nTeam name: RescueGrid A.I.D.A.\nTeam leader name: [Your Name]\nProblem Statement: {option3}"
        for p in shape.text_frame.paragraphs:
            p.font.size = Pt(16)

ppt.save(ppt_path)
print("Updated Slide 2 Problem Statement successfully.")
